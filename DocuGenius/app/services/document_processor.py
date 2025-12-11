"""Document processing service with OCR and AI classification"""

from fastapi import UploadFile
import pytesseract
from PIL import Image
import PyPDF2
import docx
import io
import os
import json
from typing import Dict, Any
from loguru import logger
import aiofiles
import httpx

from app.config import settings


class DocumentProcessor:
    """Process documents with OCR and AI classification"""
    
    def __init__(self):
        # Set Tesseract path if specified
        if settings.TESSERACT_CMD != "tesseract":
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD
    
    async def process_document(self, file: UploadFile) -> Dict[str, Any]:
        """
        Process uploaded document:
        1. Extract text (with OCR if needed)
        2. Classify using AI
        3. Return structured data
        """
        content = await file.read()
        file_type = file.filename.split(".")[-1].lower()
        
        # Extract text
        text = await self._extract_text(content, file_type)
        
        # Classify with AI
        classification = await self._classify_document(file.filename, text)
        
        return {
            "filename": file.filename,
            "file_type": file_type,
            "file_size": len(content),
            "content": text,
            "classification": classification,
            "status": "processed"
        }
    
    async def _extract_text(self, content: bytes, file_type: str) -> str:
        """Extract text from various file types"""
        try:
            if file_type == "pdf":
                return await self._extract_from_pdf(content)
            elif file_type == "docx":
                return await self._extract_from_docx(content)
            elif file_type == "txt":
                return content.decode("utf-8")
            elif file_type in ["png", "jpg", "jpeg"]:
                return await self._extract_from_image(content)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")
        except Exception as e:
            logger.error(f"Text extraction failed: {e}")
            return f"[Error extracting text: {str(e)}]"
    
    async def _extract_from_pdf(self, content: bytes) -> str:
        """Extract text from PDF"""
        try:
            pdf_file = io.BytesIO(content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text = ""
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            
            # If no text extracted, try OCR
            if not text.strip():
                logger.info("No text in PDF, attempting OCR...")
                # Note: Would need pdf2image here for full OCR support
                text = "[PDF requires OCR - install pdf2image]"
            
            return text.strip()
        except Exception as e:
            logger.error(f"PDF extraction failed: {e}")
            return f"[Error reading PDF: {str(e)}]"
    
    async def _extract_from_docx(self, content: bytes) -> str:
        """Extract text from DOCX"""
        try:
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
            
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            return text.strip()
        except Exception as e:
            logger.error(f"DOCX extraction failed: {e}")
            return f"[Error reading DOCX: {str(e)}]"
    
    async def _extract_from_image(self, content: bytes) -> str:
        """Extract text from image using Tesseract OCR"""
        try:
            image = Image.open(io.BytesIO(content))
            text = pytesseract.image_to_string(image)
            return text.strip()
        except Exception as e:
            logger.error(f"OCR failed: {e}")
            return f"[Error in OCR: {str(e)}]"
    
    async def _classify_document(self, filename: str, text: str) -> Dict[str, Any]:
        """Classify document using Ollama LLM"""
        try:
            # Limit text length for efficiency
            text_sample = text[:1000] if len(text) > 1000 else text
            
            prompt = f"""Classify this document and extract key information.

Filename: {filename}
Content: {text_sample}

Return ONLY a valid JSON object with this exact structure:
{{
  "country": ["HU"],
  "technology": ["AI"],
  "product": ["ProductA"],
  "document_type": "marketing",
  "confidence": 0.85
}}

Rules:
- country: List of country codes (HU, USA, UK, DE, etc.) or ["Other"]
- technology: List of technologies (AI, Cloud, Blockchain, etc.) or ["Other"]
- product: List of products mentioned or ["Other"]
- document_type: One of: marketing, technical, legal, financial, research, other
- confidence: Float between 0.0 and 1.0

Return ONLY the JSON, no other text."""

            # Try to call Ollama
            async with httpx.AsyncClient(timeout=30.0) as client:
                response = await client.post(
                    f"{settings.OLLAMA_BASE_URL}/api/generate",
                    json={
                        "model": settings.OLLAMA_MODEL,
                        "prompt": prompt,
                        "stream": False,
                        "format": "json"
                    }
                )
                
                if response.status_code == 200:
                    result = response.json()
                    response_text = result.get("response", "{}")
                    
                    # Parse JSON from response
                    try:
                        classification = json.loads(response_text)
                        
                        # Validate and fix structure
                        classification = self._validate_classification(classification)
                        
                        logger.info(f"AI classification: {classification['document_type']} "
                                  f"(confidence: {classification['confidence']:.2f})")
                        
                        return classification
                    except json.JSONDecodeError as e:
                        logger.warning(f"Failed to parse AI response: {e}")
                        return self._default_classification(filename, text_sample)
                else:
                    logger.warning(f"Ollama request failed: {response.status_code}")
                    return self._default_classification(filename, text_sample)
                    
        except httpx.ConnectError:
            logger.warning("Ollama not available, using rule-based classification")
            return self._default_classification(filename, text_sample)
        except Exception as e:
            logger.error(f"Classification failed: {e}")
            return self._default_classification(filename, text_sample)
    
    def _validate_classification(self, classification: Dict[str, Any]) -> Dict[str, Any]:
        """Validate and fix classification structure"""
        return {
            "country": classification.get("country", ["Other"]) if isinstance(classification.get("country"), list) else ["Other"],
            "technology": classification.get("technology", ["Other"]) if isinstance(classification.get("technology"), list) else ["Other"],
            "product": classification.get("product", ["Other"]) if isinstance(classification.get("product"), list) else ["Other"],
            "document_type": classification.get("document_type", "other"),
            "confidence": float(classification.get("confidence", 0.5))
        }
    
    def _default_classification(self, filename: str, text: str) -> Dict[str, Any]:
        """Rule-based classification when AI is not available"""
        filename_lower = filename.lower()
        text_lower = text.lower()
        
        # Detect document type from filename and content
        doc_type = "other"
        confidence = 0.6
        
        if any(word in filename_lower or word in text_lower for word in ["marketing", "campaign", "promotion"]):
            doc_type = "marketing"
        elif any(word in filename_lower or word in text_lower for word in ["technical", "specification", "architecture"]):
            doc_type = "technical"
        elif any(word in filename_lower or word in text_lower for word in ["legal", "contract", "agreement"]):
            doc_type = "legal"
        elif any(word in filename_lower or word in text_lower for word in ["financial", "budget", "revenue"]):
            doc_type = "financial"
        elif any(word in filename_lower or word in text_lower for word in ["research", "study", "analysis"]):
            doc_type = "research"
        
        # Detect countries
        countries = []
        country_keywords = {
            "HU": ["hungary", "hungarian", "magyarország"],
            "USA": ["usa", "united states", "america"],
            "UK": ["uk", "united kingdom", "britain"],
            "DE": ["germany", "german", "deutschland"]
        }
        
        for code, keywords in country_keywords.items():
            if any(kw in text_lower for kw in keywords):
                countries.append(code)
        
        if not countries:
            countries = ["Other"]
        
        # Detect technologies
        technologies = []
        tech_keywords = ["ai", "artificial intelligence", "machine learning", "cloud", "blockchain", "iot"]
        
        for tech in tech_keywords:
            if tech in text_lower:
                technologies.append(tech.upper() if len(tech) <= 3 else tech.title())
        
        if not technologies:
            technologies = ["Other"]
        
        return {
            "country": countries,
            "technology": technologies,
            "product": ["Other"],
            "document_type": doc_type,
            "confidence": confidence
        }
    
    async def check_ollama(self) -> bool:
        """Check if Ollama is available"""
        try:
            async with httpx.AsyncClient(timeout=5.0) as client:
                response = await client.get(f"{settings.OLLAMA_BASE_URL}/api/tags")
                return response.status_code == 200
        except:
            return False

